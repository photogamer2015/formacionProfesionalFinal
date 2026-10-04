"""Registro Estudiantil: las matrículas como en el Excel del equipo.

- Una hoja por día, según la fecha de matrícula. La hoja de hoy se cierra sola
  a medianoche (hora de Ecuador) y al día siguiente empieza una hoja nueva
  desde cero. Las hojas no se pierden: muestran las matrículas de ese día,
  también las que pasaron al archivo por un cierre de curso (esas solo se
  consultan).
- «Ir a rango de fecha» junta las hojas de un rango en una sola tabla, de lo
  más reciente a lo más antiguo.
- Filtros de «Lista de Matriculados» y de la Hoja de recaudación (menos la
  fecha), exportación a Excel e impresión en A4 horizontal.
- Cada fila tiene el color de la asesora que registró la matrícula: el que
  eligió el administrador en el admin de Django (ver colores_registro.py).
- Las asesoras consultan; solo el administrador edita, celda por celda (ver
  forms_registro_estudiantil.py para las reglas).
"""
from collections import Counter
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from urllib.parse import urlencode

from django.contrib.auth.models import User
from django.db import transaction
from django.db.models import Count, Prefetch, Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import redirect, render
from django.urls import reverse
from django.utils import dateformat, timezone
from django.utils.dateparse import parse_date
from django.views.decorators.http import require_POST

from .busqueda import filtrar_queryset_busqueda
from .colores_registro import Paleta
from .forms_registro_estudiantil import (
    RegistroEstudiantilForm, bloque_pago_inicial, cursos_editables,
    jornadas_editables, nombre_usuario, pago_inicial_editable,
    resumen_pago_inicial, usuarios_editables,
)
from .models import (
    Abono, AbonoArchivado, BANCOS_PAGO, BANCOS_POR_METODO, Curso,
    JornadaCurso, METODOS_PAGO, MODALIDADES, Matricula, MatriculaArchivada,
    Sede, TIPO_MATRICULA, TIPOS_REGISTRO, nombre_banco, nombre_metodo_pago,
)
from .permisos import es_admin, matricula_requerida
from .views import _aplicar_filtro_estado_pago_matricula


# Filas que se muestran como máximo en un rango; uno más largo pide acotar.
LIMITE_FILAS = 1500

# Nombres de los campos para los mensajes de error de la hoja.
ETIQUETAS_CAMPOS = {
    'fecha_matricula': 'Fecha',
    'registrado_por': 'Registra',
    'nombres': 'Nombres del estudiante',
    'vendedora': 'Vendedora',
    'tipo_registro': 'Central',
    'curso': 'Programa / Curso',
    'modalidad': 'Modalidad',
    'jornada': 'Jornada y fecha de inicio',
    'tipo_matricula': 'Matrícula',
    'valor_pago': 'Valor',
    'metodo_pago': 'Forma de pago',
    'banco': 'Banco',
    'numero_factura': 'No. de factura',
}


# ── Colores y textos ─────────────────────────────────────────────────

def _primer_nombre(nombre):
    partes = (nombre or '').split()
    return partes[0] if partes else ''


def nombre_corto(usuario):
    """Primer nombre, como en la hoja de Excel (o el usuario si no tiene)."""
    if not usuario:
        return ''
    return _primer_nombre(usuario.first_name) or usuario.username


def _fecha_texto(fecha):
    """«11 de octubre», como la columna Fecha de inicio del Excel."""
    return dateformat.format(fecha, r'j \d\e F') if fecha else ''


def _etiqueta_hoja(fecha, hoy):
    texto = dateformat.format(fecha, r'd \d\e F')
    return texto if fecha.year == hoy.year else f'{texto} {fecha.year}'


def _dinero(monto):
    return f'{monto:.2f}'.replace('.', ',') if monto is not None else ''


def _modalidad(modalidad, sede):
    texto = 'Virtual' if modalidad == 'online' else 'Presencial'
    return f'{texto} {sede}' if sede else texto


def _fecha(valor):
    try:
        return parse_date((valor or '').strip())
    except ValueError:
        return None


# ── Filas ────────────────────────────────────────────────────────────

def _fila_matricula(m, abonos, *, admin, paleta):
    bloque = bloque_pago_inicial(m, abonos)
    pago = resumen_pago_inicial(bloque) or {}
    jornada = m.jornada
    sede = jornada.sede_nombre if jornada else ''
    fila = {
        'clave': f'm{m.pk}',
        'pk': m.pk,
        'archivada': False,
        'fecha': m.fecha_matricula.isoformat(),
        'color': paleta.de_usuario(m.registrado_por),
        'retiro': m.estado == 'retiro_voluntario',
        'etiquetas': {
            'fecha': f'{m.fecha_matricula:%d/%m/%Y}',
            'registra': nombre_corto(m.registrado_por),
            'registra_completo': nombre_usuario(m.registrado_por),
            'nombres': m.estudiante.nombres,
            'vendedora': nombre_corto(m.vendedora),
            'vendedora_completo': nombre_usuario(m.vendedora),
            'central': m.get_tipo_registro_display() if m.tipo_registro else '',
            'curso': m.curso.nombre,
            'modalidad': _modalidad(m.modalidad, sede),
            'dias': jornada.descripcion_legible if jornada else '',
            'inicio': _fecha_texto(jornada.fecha_inicio) if jornada else '',
            'tipo': m.get_tipo_matricula_display(),
            'valor': _dinero(pago.get('monto')),
            'forma': pago.get('forma', ''),
            'banco': pago.get('banco', ''),
            'factura': m.numero_factura,
        },
    }
    if not admin:
        return fila

    editable = pago_inicial_editable(bloque)
    valores = {
        'fecha_matricula': m.fecha_matricula.isoformat(),
        'registrado_por': str(m.registrado_por_id or ''),
        'nombres': ' '.join((m.estudiante.nombres or '').split()),
        'vendedora': str(m.vendedora_id or ''),
        'tipo_registro': m.tipo_registro,
        'curso': str(m.curso_id),
        'modalidad': m.modalidad,
        # «sede» y «dias» solo filtran las opciones en pantalla.
        'sede': sede,
        'dias': jornada.descripcion_legible if jornada else '',
        'jornada': str(m.jornada_id or ''),
        'tipo_matricula': m.tipo_matricula,
        'numero_factura': m.numero_factura,
    }
    if editable:
        abono = bloque[0]
        valores.update(
            valor_pago=f'{abono.monto:.2f}', metodo_pago=abono.metodo,
            banco=abono.banco,
        )
    if editable:
        enlace_pago = ''
    elif bloque:
        # Pago mixto o repartido en varios abonos.
        enlace_pago = (
            reverse('academia:matricula_editar', kwargs={
                'modalidad': m.modalidad, 'pk': m.pk,
            }) + '?editar_pago=1'
        )
    else:
        enlace_pago = reverse('academia:matricula_abonos', args=[m.pk])
    fila.update(
        version=m.actualizado.isoformat(),
        valores=valores,
        # Para avisar en la confirmación si el valor del curso cambia.
        valor_curso=f'{m.valor_curso or 0:.2f}',
        descuento=f'{m.descuento or 0:.2f}',
        pago_editable=editable,
        enlace_pago=enlace_pago,
        # Lo que ya tiene valor se puede cambiar, pero no dejar vacío.
        obligatorios=[
            campo for campo, valor in (
                ('registrado_por', m.registrado_por_id),
                ('vendedora', m.vendedora_id),
                ('tipo_registro', m.tipo_registro),
            ) if valor
        ],
    )
    return fila


def _fila_archivada(a, abonos, paleta):
    bloque = []
    for abono in abonos:
        if abono.fecha != a.fecha_matricula:
            break
        bloque.append(abono)
    metodos, bancos = [], []
    for abono in bloque:
        for etiqueta, lista in (
            (abono.metodo_label or nombre_metodo_pago(abono.metodo), metodos),
            (abono.banco_label or nombre_banco(abono.banco), bancos),
        ):
            if etiqueta and etiqueta not in lista:
                lista.append(etiqueta)
    tipos_registro = dict(TIPOS_REGISTRO)
    return {
        'clave': f'a{a.pk}',
        'pk': None,
        'archivada': True,
        'fecha': a.fecha_matricula.isoformat(),
        'color': paleta.de_nombre(a.registrado_por_nombre),
        'retiro': a.estado == 'retiro_voluntario',
        'etiquetas': {
            'fecha': f'{a.fecha_matricula:%d/%m/%Y}',
            'registra': _primer_nombre(a.registrado_por_nombre),
            'registra_completo': a.registrado_por_nombre,
            'nombres': a.nombres,
            'vendedora': _primer_nombre(a.vendedora_nombre),
            'vendedora_completo': a.vendedora_nombre,
            'central': tipos_registro.get(a.tipo_registro, a.tipo_registro),
            'curso': a.curso_nombre,
            'modalidad': _modalidad(a.modalidad, a.sede),
            'dias': a.jornada_descripcion,
            'inicio': _fecha_texto(a.jornada_fecha_inicio),
            'tipo': dict(TIPO_MATRICULA).get(a.tipo_matricula, a.tipo_matricula),
            'valor': _dinero(sum((ab.monto for ab in bloque), 0)) if bloque else '',
            'forma': ' + '.join(metodos),
            'banco': ' + '.join(bancos),
            'factura': a.numero_factura,
        },
    }


def _con_abonos(qs):
    return qs.select_related(
        'estudiante', 'curso', 'jornada', 'jornada__sede',
        'registrado_por', 'vendedora',
    ).prefetch_related(
        Prefetch(
            'abonos', queryset=Abono.objects.order_by('creado', 'id'),
            to_attr='abonos_registro',
        )
    )


def _archivadas():
    """Matrículas que pasaron al archivo por un cierre (y ya no están vivas)."""
    return MatriculaArchivada.objects.exclude(
        matricula_original_id__in=Matricula.objects.values('pk'),
    )


def _filas(matriculas, archivadas, *, admin, reciente_primero, limite=None):
    """Une matrículas vivas y archivadas en el orden de la hoja y las numera.

    Por día van en el orden en que se registraron; en un rango, de lo más
    reciente a lo más antiguo.
    """
    matriculas = _con_abonos(matriculas)
    archivadas = archivadas.prefetch_related(
        Prefetch(
            'abonos_archivados',
            queryset=AbonoArchivado.objects.order_by('creado_original', 'id'),
            to_attr='abonos_registro',
        )
    )
    if limite:
        signo = '-' if reciente_primero else ''
        matriculas = matriculas.order_by(
            f'{signo}fecha_matricula', f'{signo}creado', f'{signo}pk',
        )[:limite]
        archivadas = archivadas.order_by(
            f'{signo}fecha_matricula', f'{signo}creado_original', f'{signo}pk',
        )[:limite]
    paleta = Paleta()
    orden = [
        ((m.fecha_matricula, m.creado, m.pk),
         _fila_matricula(m, m.abonos_registro, admin=admin, paleta=paleta))
        for m in matriculas
    ] + [
        ((a.fecha_matricula, a.creado_original or a.archivado_en, a.pk),
         _fila_archivada(a, a.abonos_registro, paleta))
        for a in archivadas
    ]
    orden.sort(key=lambda item: item[0], reverse=reciente_primero)
    filas = [fila for _, fila in orden][:limite]
    for numero, fila in enumerate(filas, 1):
        fila['numero'] = numero
    return filas


def _leyenda(filas):
    """Cuántas matrículas registró cada asesora, con su color."""
    conteo = Counter(
        (f['etiquetas']['registra_completo'] or 'Sin registrar', f['color'])
        for f in filas
    )
    return [
        {'nombre': nombre, 'color': color, 'total': total}
        for (nombre, color), total in sorted(
            conteo.items(), key=lambda item: (-item[1], item[0][0].casefold()),
        )
    ]


def _hojas(hoy, seleccionada):
    """Pestañas: un día por hoja, de la más reciente a la más antigua."""
    conteo = Counter()
    for qs in (Matricula.objects.all(), _archivadas()):
        for fecha, total in (
            qs.order_by().values_list('fecha_matricula').annotate(n=Count('pk'))
        ):
            conteo[fecha] += total
    conteo.setdefault(hoy, 0)
    conteo.setdefault(seleccionada, 0)
    hojas = [
        {
            'fecha': fecha,
            'iso': fecha.isoformat(),
            'etiqueta': _etiqueta_hoja(fecha, hoy),
            'total': total,
            'activa': fecha == seleccionada,
            'hoy': fecha == hoy,
        }
        for fecha, total in sorted(conteo.items(), reverse=True)
    ]
    meses = []
    for hoja in hojas:
        mes = hoja['fecha'].replace(day=1)
        if not meses or meses[-1]['mes'] != mes:
            meses.append({'mes': mes, 'hojas': []})
        meses[-1]['hojas'].append(hoja)
    return hojas, meses


def _datos_edicion(filas, vista):
    """Opciones de las celdas editables (solo administrador)."""
    vivas = Matricula.objects.filter(
        pk__in=[f['pk'] for f in filas if not f['archivada']],
    )
    usuarios_actuales, cursos_actuales, jornadas_actuales = set(), set(), set()
    for registra, vendedora, curso, jornada in vivas.values_list(
        'registrado_por_id', 'vendedora_id', 'curso_id', 'jornada_id',
    ):
        usuarios_actuales.update((registra, vendedora))
        cursos_actuales.add(curso)
        jornadas_actuales.add(jornada)

    usuarios_generales = set(usuarios_editables().values_list('pk', flat=True))
    cursos_generales = set(cursos_editables().values_list('pk', flat=True))
    cursos = cursos_editables(*cursos_actuales)
    return {
        'vista': vista,
        'url_guardar': reverse('academia:registro_estudiantil_guardar', args=[0]),
        'usuarios': [
            {'id': str(u.pk), 'nombre': nombre_usuario(u),
             'general': u.pk in usuarios_generales}
            for u in usuarios_editables(*usuarios_actuales)
        ],
        'cursos': [
            {'id': str(c.pk),
             'nombre': c.nombre + (' (Ciclo Corto)' if c.es_ciclo_corto else ''),
             'general': c.pk in cursos_generales,
             'valor_presencial': f'{c.valor_presencial or 0:.2f}',
             'valor_online': f'{c.valor_online or 0:.2f}'}
            for c in cursos
        ],
        'jornadas': [
            {
                'id': str(j.pk), 'curso': str(j.curso_id),
                'modalidad': j.modalidad, 'sede': j.sede_nombre,
                'dias': j.descripcion_legible,
                'fecha': j.fecha_inicio.isoformat() if j.fecha_inicio else '',
                'fecha_texto': _fecha_texto(j.fecha_inicio) or 'Sin fecha',
                'horario': (
                    f'{j.hora_inicio:%H:%M}–{j.hora_fin:%H:%M}'
                    if j.hora_inicio and j.hora_fin else ''
                ),
                'activo': j.activo,
            }
            for j in jornadas_editables(cursos, *jornadas_actuales)
        ],
        'centrales': TIPOS_REGISTRO,
        'tipos': TIPO_MATRICULA,
        'metodos': METODOS_PAGO,
        'bancos': dict(BANCOS_PAGO),
        'bancos_por_metodo': BANCOS_POR_METODO,
    }


def _contexto_tabla(request, filas, vista):
    admin = es_admin(request.user)
    contexto = {
        'filas': filas,
        'leyenda': _leyenda(filas),
        'puede_editar': admin,
    }
    if admin:
        contexto['datos_edicion'] = _datos_edicion(filas, vista)
        contexto['filas_edicion'] = {
            f['clave']: f for f in filas if not f['archivada']
        }
    return contexto


# ── Filtros (los de «Lista de Matriculados» y la Hoja de recaudación) ──

FILTROS_REGISTRO = (
    'q', 'curso', 'jornada', 'descuento', 'estado_pago', 'modalidad',
    'ciudad', 'sede', 'registrador',
)
# Los mismos campos que el buscador de «Lista de Matriculados».
BUSQUEDA_VIVAS = [
    'estudiante__cedula', 'estudiante__nombres', 'estudiante__correo',
    'estudiante__celular', 'curso__nombre', 'fact_cedula', 'fact_nombres',
]
BUSQUEDA_ARCHIVADAS = [
    'cedula', 'nombres', 'correo', 'celular', 'curso_nombre',
    'fact_cedula', 'fact_nombres',
]
ESTADOS_PAGO = {'pendiente': 'Saldo pendiente', 'pagado': 'Pagado'}
DESCUENTOS = {'si': 'Con descuento', 'no': 'Sin descuento'}


def _leer_filtros(request):
    """Filtros de la consulta, ya validados (lo que no sirve se ignora)."""
    filtros = {
        clave: (request.GET.get(clave) or '').strip()
        for clave in FILTROS_REGISTRO
    }
    for clave in ('curso', 'jornada', 'registrador'):
        if not filtros[clave].isdigit():
            filtros[clave] = ''
    if not filtros['curso']:
        # La jornada se elige dentro de un curso, como en la Hoja de recaudación.
        filtros['jornada'] = ''
    if filtros['descuento'] not in DESCUENTOS:
        filtros['descuento'] = ''
    if filtros['estado_pago'] not in ESTADOS_PAGO:
        filtros['estado_pago'] = ''
    if filtros['modalidad'] not in dict(MODALIDADES):
        filtros['modalidad'] = ''
    sede = filtros['sede']
    if not (
        (sede.startswith('sede:') and sede[5:].isdigit())
        or (sede.startswith('ciudad:') and sede[7:])
    ):
        filtros['sede'] = ''
    return filtros


def _aplicar_filtros(matriculas, archivadas, filtros):
    """Aplica los filtros a las matrículas vivas y a las archivadas."""
    if filtros['curso']:
        matriculas = matriculas.filter(curso_id=int(filtros['curso']))
        archivadas = archivadas.filter(curso_id=int(filtros['curso']))
    if filtros['jornada']:
        matriculas = matriculas.filter(jornada_id=int(filtros['jornada']))
        archivadas = archivadas.filter(jornada_id=int(filtros['jornada']))
    if filtros['descuento'] == 'si':
        matriculas = matriculas.filter(descuento__gt=0)
        archivadas = archivadas.filter(descuento__gt=0)
    elif filtros['descuento'] == 'no':
        matriculas = matriculas.filter(descuento=0)
        archivadas = archivadas.filter(descuento=0)
    if filtros['estado_pago']:
        matriculas = _aplicar_filtro_estado_pago_matricula(
            matriculas, filtros['estado_pago'],
        )
        archivadas = archivadas.filter(estado_pago__in=(
            ('Parcial', 'Pendiente') if filtros['estado_pago'] == 'pendiente'
            else ('Pagado',)
        ))
    if filtros['modalidad']:
        matriculas = matriculas.filter(modalidad=filtros['modalidad'])
        archivadas = archivadas.filter(modalidad=filtros['modalidad'])
    if filtros['ciudad']:
        matriculas = matriculas.filter(jornada__ciudad__iexact=filtros['ciudad'])
        archivadas = archivadas.filter(sede__iexact=filtros['ciudad'])
    if filtros['sede'].startswith('sede:'):
        sede_id = int(filtros['sede'][5:])
        nombre_sede = (
            Sede.objects.filter(pk=sede_id).values_list('nombre', flat=True).first()
        )
        matriculas = matriculas.filter(jornada__sede_id=sede_id)
        archivadas = (
            archivadas.filter(sede__iexact=nombre_sede) if nombre_sede
            else archivadas.none()
        )
    elif filtros['sede'].startswith('ciudad:'):
        matriculas = matriculas.filter(jornada__ciudad=filtros['sede'][7:])
        archivadas = archivadas.filter(sede=filtros['sede'][7:])
    if filtros['registrador']:
        usuario = User.objects.filter(pk=int(filtros['registrador'])).first()
        matriculas = matriculas.filter(registrado_por_id=int(filtros['registrador']))
        # Las archivadas guardan el nombre de quien registró.
        archivadas = (
            archivadas.filter(registrado_por_nombre=nombre_usuario(usuario))
            if usuario else archivadas.none()
        )
    if filtros['q']:
        matriculas = filtrar_queryset_busqueda(matriculas, filtros['q'], BUSQUEDA_VIVAS)
        archivadas = filtrar_queryset_busqueda(
            archivadas, filtros['q'], BUSQUEDA_ARCHIVADAS,
        )
    return matriculas, archivadas


def _etiqueta_jornada_filtro(jornada):
    texto = jornada.etiqueta
    return texto if jornada.activo else f'{texto} (desactivada)'


def _opciones_filtros(filtros):
    """Opciones de los selectores de filtro."""
    cursos = Curso.objects.filter(
        Q(pk=filtros['curso'] or None)
        | Q(activo=True) & (Q(ofrece_presencial=True) | Q(ofrece_online=True))
    ).order_by('nombre')
    jornadas = JornadaCurso.objects.filter(
        curso__in=cursos,
    ).select_related('sede').order_by('fecha_inicio', 'pk')
    sedes = [
        {'value': f'sede:{sede.pk}', 'label': sede.etiqueta}
        for sede in Sede.objects.filter(jornadas__isnull=False)
        .distinct().order_by('pais', 'orden', 'nombre')
    ] + [
        {'value': f'ciudad:{ciudad}', 'label': ciudad}
        for ciudad in JornadaCurso.objects.filter(sede__isnull=True)
        .exclude(ciudad='').values_list('ciudad', flat=True)
        .distinct().order_by('ciudad')
    ]
    ciudades = sorted(
        set(
            JornadaCurso.objects.filter(modalidad='presencial')
            .exclude(ciudad='').values_list('ciudad', flat=True)
        ),
        key=str.casefold,
    )
    if filtros['ciudad'] and filtros['ciudad'] not in ciudades:
        ciudades.append(filtros['ciudad'])
    return {
        'cursos': cursos,
        'jornadas_filtro': [
            {
                'id': str(j.pk), 'curso': str(j.curso_id),
                'modalidad': j.modalidad, 'ciudad': j.ciudad or '',
                'texto': _etiqueta_jornada_filtro(j),
            }
            for j in jornadas
        ],
        'sedes': sedes,
        'ciudades': ciudades,
        'registradores': User.objects.filter(is_active=True).order_by(
            'first_name', 'username',
        ),
    }


def _resumen_filtros(filtros, opciones):
    """Filtros aplicados en palabras, para el Excel y la impresión."""
    partes = []
    if filtros['q']:
        partes.append(f'Búsqueda: «{filtros["q"]}»')
    if filtros['curso']:
        curso = next(
            (c for c in opciones['cursos'] if str(c.pk) == filtros['curso']), None,
        )
        partes.append(f'Curso: {curso.nombre if curso else filtros["curso"]}')
    if filtros['jornada']:
        jornada = JornadaCurso.objects.filter(pk=int(filtros['jornada'])).first()
        if jornada:
            partes.append(f'Jornada: {_etiqueta_jornada_filtro(jornada)}')
    if filtros['descuento']:
        partes.append(DESCUENTOS[filtros['descuento']])
    if filtros['estado_pago']:
        partes.append(f'Estado de pago: {ESTADOS_PAGO[filtros["estado_pago"]]}')
    if filtros['modalidad']:
        partes.append(f'Modalidad: {dict(MODALIDADES)[filtros["modalidad"]]}')
    if filtros['ciudad']:
        partes.append(f'Ciudad: {filtros["ciudad"]}')
    if filtros['sede']:
        etiqueta = next(
            (s['label'] for s in opciones['sedes'] if s['value'] == filtros['sede']),
            filtros['sede'].split(':', 1)[-1],
        )
        partes.append(f'Sede: {etiqueta}')
    if filtros['registrador']:
        usuario = User.objects.filter(pk=int(filtros['registrador'])).first()
        partes.append(f'Registrador: {nombre_usuario(usuario) if usuario else "—"}')
    return partes


# ── Consulta común (hoja, rango, Excel e impresión) ─────────────────

def _periodo(request, modo):
    """Días que abarca la consulta: un día (hoja) o un rango de fechas."""
    if modo == 'hoja':
        hoja = _fecha(request.GET.get('hoja')) or timezone.localdate()
        return hoja, hoja
    desde = _fecha(request.GET.get('fecha_desde'))
    hasta = _fecha(request.GET.get('fecha_hasta'))
    if desde and not hasta:
        hasta = desde
    if hasta and not desde:
        desde = hasta
    if desde and hasta and desde > hasta:
        desde, hasta = hasta, desde
    return desde, hasta


def _consulta(request, modo, *, admin):
    desde, hasta = _periodo(request, modo)
    filtros = _leer_filtros(request)
    filas, total, total_sin_filtros = [], 0, 0
    if desde and hasta:
        matriculas = Matricula.objects.filter(fecha_matricula__range=(desde, hasta))
        archivadas = _archivadas().filter(fecha_matricula__range=(desde, hasta))
        total_sin_filtros = matriculas.count() + archivadas.count()
        matriculas, archivadas = _aplicar_filtros(matriculas, archivadas, filtros)
        total = matriculas.count() + archivadas.count()
        filas = _filas(
            matriculas, archivadas, admin=admin,
            reciente_primero=modo == 'rango',
            limite=LIMITE_FILAS if modo == 'rango' else None,
        )
    filtros_query = urlencode({k: v for k, v in filtros.items() if v})
    if modo == 'hoja':
        periodo_query = urlencode({'hoja': desde.isoformat()})
    elif desde:
        periodo_query = urlencode({
            'fecha_desde': desde.isoformat(), 'fecha_hasta': hasta.isoformat(),
        })
    else:
        periodo_query = ''
    return {
        'modo': modo,
        'desde': desde,
        'hasta': hasta,
        'filas': filas,
        'total': total,
        'total_sin_filtros': total_sin_filtros,
        'recortado': total > len(filas),
        'filtros': filtros,
        'hay_filtros': bool(filtros_query),
        'filtros_query': filtros_query,
        'periodo_query': periodo_query,
        'consulta_query': '&'.join(q for q in (periodo_query, filtros_query) if q),
    }


def _titulo(consulta):
    desde, hasta = consulta['desde'], consulta['hasta']
    if consulta['modo'] == 'hoja':
        return 'Hoja del ' + dateformat.format(desde, r'l d \d\e F \d\e Y')
    if desde == hasta:
        return f'Matrículas del {desde:%d/%m/%Y}'
    return f'Matrículas del {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}'


def _contexto_filtros(consulta):
    opciones = _opciones_filtros(consulta['filtros'])
    url = reverse(
        'academia:registro_estudiantil' if consulta['modo'] == 'hoja'
        else 'academia:registro_estudiantil_rango'
    )
    return {
        **opciones,
        # «Limpiar» quita los filtros pero se queda en el mismo día o rango.
        'url_limpiar': f'{url}?{consulta["periodo_query"]}' if consulta['periodo_query'] else url,
        'puede_exportar': bool(consulta['desde']),
        'filtros': consulta['filtros'],
        'hay_filtros': consulta['hay_filtros'],
        'filtros_query': consulta['filtros_query'],
        'consulta_query': consulta['consulta_query'],
        'total': consulta['total'],
        'total_sin_filtros': consulta['total_sin_filtros'],
    }


# ── Vistas ───────────────────────────────────────────────────────────

@matricula_requerida
def registro_estudiantil(request):
    """Una hoja: las matrículas de un día, en el orden en que se registraron."""
    ahora = timezone.localtime()
    hoy = ahora.date()
    consulta = _consulta(request, 'hoja', admin=es_admin(request.user))
    hoja = consulta['desde']
    hojas, meses = _hojas(hoy, hoja)
    medianoche = timezone.make_aware(
        datetime.combine(hoy + timedelta(days=1), time.min),
        timezone.get_current_timezone(),
    )
    return render(request, 'registro_estudiantil/hoja.html', {
        **_contexto_tabla(
            request, consulta['filas'], {'modo': 'hoja', 'hoja': hoja.isoformat()},
        ),
        **_contexto_filtros(consulta),
        'modo': 'hoja',
        'hoja': hoja,
        'hoy': hoy,
        'es_hoy': hoja == hoy,
        'hojas': hojas,
        'meses_hojas': meses,
        'segundos_para_cerrar': int((medianoche - ahora).total_seconds()) + 1,
    })


@matricula_requerida
def registro_estudiantil_rango(request):
    """Las hojas de un rango de fechas juntas, de lo más reciente a lo más
    antiguo. Sin rango solo se muestra el selector de fechas."""
    consulta = _consulta(request, 'rango', admin=es_admin(request.user))
    desde, hasta = consulta['desde'], consulta['hasta']
    vista = {
        'modo': 'rango',
        'desde': desde.isoformat() if desde else '',
        'hasta': hasta.isoformat() if hasta else '',
    }
    return render(request, 'registro_estudiantil/rango.html', {
        **_contexto_tabla(request, consulta['filas'], vista),
        **_contexto_filtros(consulta),
        'modo': 'rango',
        'desde': desde,
        'hasta': hasta,
        'fecha_desde': vista['desde'],
        'fecha_hasta': vista['hasta'],
        'recortado': consulta['recortado'],
        'limite': LIMITE_FILAS,
    })


@matricula_requerida
@require_POST
@transaction.atomic
def registro_estudiantil_guardar(request, pk):
    """Guarda la fila que el administrador cambió en la hoja (JSON)."""
    if not es_admin(request.user):
        return JsonResponse({
            'ok': False,
            'mensajes': ['Solo un administrador puede editar el Registro Estudiantil.'],
        }, status=403)
    matricula = Matricula.objects.select_for_update().filter(pk=pk).first()
    if matricula is None:
        return JsonResponse({
            'ok': False,
            'conflicto': True,
            'mensajes': [
                'Esta matrícula ya no existe: pudo eliminarse o pasar al '
                'archivo por un cierre de curso. Recarga la hoja.'
            ],
        }, status=404)

    form = RegistroEstudiantilForm(request.POST, matricula=matricula)
    if not form.is_valid():
        mensajes = [
            f'{ETIQUETAS_CAMPOS[campo]}: {error}'
            if campo in ETIQUETAS_CAMPOS else str(error)
            for campo, errores in form.errors.items()
            for error in errores
        ]
        return JsonResponse({
            'ok': False,
            'conflicto': form.conflicto,
            'campos': list(form.errors),
            'mensajes': mensajes,
        }, status=400)

    cambios = form.guardar(request.user)
    actualizada = _con_abonos(Matricula.objects.filter(pk=pk)).get()
    return JsonResponse({
        'ok': True,
        'cambios': cambios,
        'fila': _fila_matricula(
            actualizada, actualizada.abonos_registro, admin=True,
            paleta=Paleta(),
        ),
    })


# ── Excel e impresión ────────────────────────────────────────────────

# Columnas como en la hoja (y en el Excel del equipo), con su ancho en Excel.
COLUMNAS_EXCEL = [
    ('No.', 6), ('Fecha', 12), ('Registra', 13),
    ('Nombres del Estudiante', 34), ('Vendedora', 13), ('Central', 16),
    ('Programa / Curso', 28), ('Modalidad', 21), ('Jornada', 20),
    ('Fecha de inicio', 15), ('Matrícula', 19), ('Valor', 10),
    ('Forma de Pago', 21), ('Banco', 18), ('No. de Factura', 16),
]


def _modo_salida(request):
    """Excel e impresión sirven a la hoja del día y al rango de fechas."""
    if 'fecha_desde' in request.GET or 'fecha_hasta' in request.GET:
        return 'rango'
    return 'hoja'


def _monto(fila):
    try:
        return Decimal(fila['etiquetas']['valor'].replace(',', '.'))
    except (InvalidOperation, AttributeError):
        return None


def _nombre_con_marcas(fila):
    nombre = fila['etiquetas']['nombres']
    if fila['retiro']:
        nombre += ' (retiro)'
    if fila['archivada']:
        nombre += ' (archivada)'
    return nombre


def _url_volver(consulta):
    nombre = (
        'academia:registro_estudiantil' if consulta['modo'] == 'hoja'
        else 'academia:registro_estudiantil_rango'
    )
    url = reverse(nombre)
    return f'{url}?{consulta["consulta_query"]}' if consulta['consulta_query'] else url


@matricula_requerida
def registro_estudiantil_excel(request):
    """La hoja (o el rango) con sus filtros en un .xlsx con los mismos
    colores, listo para imprimir en A4 horizontal."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    consulta = _consulta(request, _modo_salida(request), admin=False)
    if not consulta['desde']:
        return redirect('academia:registro_estudiantil_rango')
    filas = consulta['filas']
    resumen = _resumen_filtros(consulta['filtros'], _opciones_filtros(consulta['filtros']))
    desde, hasta = consulta['desde'], consulta['hasta']

    wb = Workbook()
    ws = wb.active
    if consulta['modo'] == 'hoja':
        ws.title = f'{desde:%d-%m-%Y}'
        archivo = f'registro_estudiantil_{desde:%Y-%m-%d}.xlsx'
    else:
        ws.title = f'{desde:%d-%m-%Y} al {hasta:%d-%m-%Y}'[:31]
        archivo = f'registro_estudiantil_{desde:%Y-%m-%d}_al_{hasta:%Y-%m-%d}.xlsx'

    columnas = len(COLUMNAS_EXCEL)
    ultima = get_column_letter(columnas)
    linea_blanca = Side(style='thin', color='FFFFFF')
    borde = Border(left=linea_blanca, right=linea_blanca, top=linea_blanca, bottom=linea_blanca)

    ws.merge_cells(f'A1:{ultima}1')
    ws['A1'] = f'Registro Estudiantil · {_titulo(consulta)}'
    ws['A1'].font = Font(bold=True, size=14, color='1B2A4E')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 24
    ws.merge_cells(f'A2:{ultima}2')
    ws['A2'] = (
        f'{len(filas)} registro{"" if len(filas) == 1 else "s"} · '
        + ('Filtros: ' + '; '.join(resumen) if resumen else 'Sin filtros')
        + f' · Generado el {timezone.localtime():%d/%m/%Y %H:%M}'
    )
    ws['A2'].font = Font(italic=True, size=10, color='4B5563')
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    encabezado_relleno = PatternFill('solid', fgColor='FF2FD34F')
    encabezado_letra = Font(name='Georgia', bold=True, italic=True, size=11, color='0B2A10')
    for indice, (texto, ancho) in enumerate(COLUMNAS_EXCEL, start=1):
        celda = ws.cell(row=3, column=indice, value=texto)
        celda.fill = encabezado_relleno
        celda.font = encabezado_letra
        celda.border = borde
        celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(indice)].width = ancho
    ws.row_dimensions[3].height = 24

    total_valor = Decimal('0.00')
    for numero_fila, fila in enumerate(filas, start=4):
        e = fila['etiquetas']
        monto = _monto(fila)
        if monto is not None:
            total_valor += monto
        valores = [
            fila['numero'], date.fromisoformat(fila['fecha']), e['registra'],
            _nombre_con_marcas(fila), e['vendedora'], e['central'], e['curso'],
            e['modalidad'], e['dias'], e['inicio'], e['tipo'],
            float(monto) if monto is not None else None, e['forma'], e['banco'],
            e['factura'],
        ]
        relleno = PatternFill('solid', fgColor='FF' + fila['color'].lstrip('#').upper())
        for indice, valor in enumerate(valores, start=1):
            celda = ws.cell(
                row=numero_fila, column=indice,
                value=None if valor in ('', None) else valor,
            )
            celda.fill = relleno
            celda.border = borde
            celda.font = Font(
                name='Georgia', italic=True, bold=indice in (1, 4), size=11,
                color='111111',
            )
            celda.alignment = Alignment(
                vertical='center', wrap_text=True,
                horizontal='right' if indice in (1, 12) else None,
            )
        ws.cell(row=numero_fila, column=2).number_format = 'dd/mm/yyyy'
        ws.cell(row=numero_fila, column=12).number_format = '#,##0.00'
        ws.cell(row=numero_fila, column=15).number_format = '@'

    fila_total = len(filas) + 4
    total_letra = Font(bold=True, color='1B2A4E', size=11)
    total_relleno = PatternFill('solid', fgColor='FFFFF8E1')
    for indice in range(1, columnas + 1):
        celda = ws.cell(row=fila_total, column=indice)
        celda.fill = total_relleno
        celda.font = total_letra
    ws.cell(row=fila_total, column=1, value='Total')
    ws.cell(
        row=fila_total, column=4,
        value=f'{len(filas)} registro{"" if len(filas) == 1 else "s"}',
    )
    celda_total = ws.cell(row=fila_total, column=12, value=float(total_valor))
    celda_total.number_format = '#,##0.00'
    celda_total.alignment = Alignment(horizontal='right')

    # Excel: encabezado fijo y filtros en cada columna.
    ws.freeze_panes = 'A4'
    ws.auto_filter.ref = f'A3:{ultima}{max(fila_total - 1, 3)}'
    # Impresión: A4 horizontal, todas las columnas a lo ancho de la hoja.
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(
        left=0.25, right=0.25, top=0.4, bottom=0.45, header=0.2, footer=0.2,
    )
    ws.print_title_rows = '3:3'
    ws.print_area = f'A1:{ultima}{fila_total}'
    ws.oddFooter.left.text = 'Formación Profesional EC · Registro Estudiantil'
    ws.oddFooter.left.size = 9
    ws.oddFooter.right.text = 'Página &P de &N'
    ws.oddFooter.right.size = 9

    salida = BytesIO()
    wb.save(salida)
    respuesta = HttpResponse(
        salida.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    respuesta['Content-Disposition'] = f'attachment; filename="{archivo}"'
    return respuesta


@matricula_requerida
def registro_estudiantil_imprimir(request):
    """Página para imprimir la hoja (o el rango) en A4 horizontal."""
    consulta = _consulta(request, _modo_salida(request), admin=False)
    if not consulta['desde']:
        return redirect('academia:registro_estudiantil_rango')
    filas = consulta['filas']
    for fila in filas:
        fila['nombre_impresion'] = _nombre_con_marcas(fila)
    montos = [m for m in (_monto(f) for f in filas) if m is not None]
    return render(request, 'registro_estudiantil/imprimir.html', {
        'titulo': _titulo(consulta),
        'filas': filas,
        'leyenda': _leyenda(filas),
        'resumen_filtros': _resumen_filtros(
            consulta['filtros'], _opciones_filtros(consulta['filtros']),
        ),
        'total_valor': f'{sum(montos, Decimal("0.00")):.2f}'.replace('.', ','),
        'recortado': consulta['recortado'],
        'total': consulta['total'],
        'limite': LIMITE_FILAS,
        'generado': timezone.localtime(),
        'url_volver': _url_volver(consulta),
    })
