from datetime import date, datetime, timedelta
from decimal import Decimal

from django.contrib.admin.models import LogEntry
from django.contrib.auth.models import Group, User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from .models import (
    Abono, AbonoArchivado, CierreCurso, Comprobante, Curso, Estudiante,
    JornadaCurso, Matricula, MatriculaArchivada, PerfilUsuario,
    RecuperacionPendiente, Sede,
)
from .colores_registro import (
    AUTOMATICOS, COLOR_SIN_REGISTRO, HEX, Paleta, color_de_usuario,
)


class RegistroBase(TestCase):
    def setUp(self):
        asesores = Group.objects.create(name='Asesores')
        self.admin = User.objects.create_superuser('admin_reg', password='x')
        self.shirley = User.objects.create_user('shirley.m', first_name='Shirley', last_name='Mora')
        self.kim = User.objects.create_user('kim')
        self.glenda = User.objects.create_user('glenda', first_name='Glenda')
        for u in (self.shirley, self.kim, self.glenda):
            u.groups.add(asesores)

        self.gye = Sede.objects.create(nombre='Guayaquil')
        self.curso = Curso.objects.create(
            nombre='Tributación Contable', valor_presencial=Decimal('90.00'),
            valor_online=Decimal('60.00'), ofrece_online=True,
            numero_modulos=4, numero_modulos_online=2,
        )
        self.otro_curso = Curso.objects.create(
            nombre='Excel Gerencial', valor_presencial=Decimal('120.00'),
            numero_modulos=3,
        )
        self.j_dom = JornadaCurso.objects.create(
            curso=self.curso, modalidad='presencial',
            descripcion='domingos_intensivos', fecha_inicio=date(2026, 10, 11),
            sede=self.gye,
        )
        self.j_sab = JornadaCurso.objects.create(
            curso=self.curso, modalidad='presencial',
            descripcion='sabados_intensivos', fecha_inicio=date(2026, 10, 10),
            sede=self.gye,
        )
        self.j_online = JornadaCurso.objects.create(
            curso=self.curso, modalidad='online',
            descripcion='mar_mie_jue', fecha_inicio=date(2026, 10, 20),
            ciudad='Ecuador',
        )
        self.j_excel = JornadaCurso.objects.create(
            curso=self.otro_curso, modalidad='presencial',
            descripcion='lun_mie_vie', fecha_inicio=date(2026, 10, 5),
            sede=self.gye,
        )

        self.estudiante = Estudiante.objects.create(
            cedula='0912345678', nombres='Chango Zambrano Daniel',
            correo='daniel@example.com',
        )
        self.m = Matricula.objects.create(
            estudiante=self.estudiante, curso=self.curso, jornada=self.j_dom,
            modalidad='presencial', tipo_matricula='reserva_abono',
            forma_pago='abono', fecha_matricula=date(2026, 10, 1),
            valor_curso=Decimal('90.00'), registrado_por=self.shirley,
            vendedora=self.glenda, tipo_registro='central_ia',
        )
        self.reserva = Abono.objects.create(
            matricula=self.m, fecha=date(2026, 10, 1), monto=Decimal('10.00'),
            metodo='transferencia', banco='pichincha', tipo_pago='abono',
        )
        self.posterior = Abono.objects.create(
            matricula=self.m, fecha=date(2026, 10, 15), monto=Decimal('20.00'),
            metodo='efectivo', tipo_pago='por_modulo', numero_modulo=1,
        )
        self.m.refresh_from_db()

        self.url_hoja = reverse('academia:registro_estudiantil')
        self.url_rango = reverse('academia:registro_estudiantil_rango')
        self.url_guardar = reverse('academia:registro_estudiantil_guardar', args=[self.m.pk])

    def matricula(self, nombre, fecha, registra=None, cedula=None, **extra):
        return Matricula.objects.create(
            estudiante=Estudiante.objects.create(
                cedula=cedula or f'09{abs(hash(nombre)) % 10**8:08d}', nombres=nombre,
            ),
            curso=self.curso, jornada=self.j_sab, modalidad='presencial',
            tipo_matricula='programa_completo', fecha_matricula=fecha,
            valor_curso=Decimal('90.00'), registrado_por=registra,
            vendedora=registra, **extra,
        )

    def archivada(self, fecha, original_id=None, nombre='Ana Archivada'):
        cierre = CierreCurso.objects.create(curso_nombre='Tributación Contable')
        archivada = MatriculaArchivada.objects.create(
            cierre=cierre, matricula_original_id=original_id, cedula='0999999999',
            nombres=nombre, curso_nombre='Tributación Contable',
            modalidad='presencial', fecha_matricula=fecha, estado_pago='Pagado',
            tipo_matricula='reserva_abono', tipo_registro='central_1',
            registrado_por_nombre='Shirley Mora', vendedora_nombre='Glenda',
            jornada_descripcion='Sábados Intensivos', sede='Guayaquil',
            jornada_fecha_inicio=date(2026, 10, 10), numero_factura='10443',
            creado_original=timezone.make_aware(datetime(2026, 10, 1, 9, 0)),
        )
        AbonoArchivado.objects.create(
            matricula_archivada=archivada, cierre=cierre, fecha=fecha,
            monto=Decimal('10.00'), metodo='deposito', metodo_label='Depósito',
            banco='guayaquil', banco_label='Guayaquil',
        )
        return archivada

    def datos(self, **cambios):
        """La fila tal como la envía la hoja, con los cambios pedidos."""
        self.m.refresh_from_db()
        datos = {
            'version': self.m.actualizado.isoformat(),
            'fecha_matricula': self.m.fecha_matricula.isoformat(),
            'registrado_por': self.m.registrado_por_id or '',
            'nombres': self.m.estudiante.nombres,
            'vendedora': self.m.vendedora_id or '',
            'tipo_registro': self.m.tipo_registro,
            'curso': self.m.curso_id,
            'modalidad': self.m.modalidad,
            'jornada': self.m.jornada_id or '',
            'tipo_matricula': self.m.tipo_matricula,
            'valor_pago': '10.00',
            'metodo_pago': 'transferencia',
            'banco': 'pichincha',
            'numero_factura': self.m.numero_factura,
        }
        datos.update(cambios)
        return datos

    def guardar(self, **cambios):
        self.client.force_login(self.admin)
        respuesta = self.client.post(self.url_guardar, self.datos(**cambios))
        return respuesta, respuesta.json()


class ColoresTests(TestCase):
    def test_sin_color_elegido_usa_los_del_excel_o_uno_automatico(self):
        paleta = Paleta()
        shirley = User(pk=1, username='s1', first_name='Shirley')
        kimberly = User(pk=2, username='kim')
        buffer = User(pk=3, username='bufer')
        melanie = User(pk=4, username='mel', first_name='Melanie Andrea')
        otra = User(pk=5, username='glenda', first_name='Glenda')
        self.assertEqual(paleta.de_usuario(shirley), HEX['azul'])
        self.assertEqual(paleta.de_usuario(kimberly), HEX['naranja'])
        self.assertEqual(paleta.de_usuario(buffer), HEX['rojo'])
        self.assertEqual(paleta.de_usuario(melanie), HEX['verde'])
        self.assertIn(paleta.de_usuario(otra), [HEX[c] for c in AUTOMATICOS])
        self.assertEqual(paleta.de_usuario(None), COLOR_SIN_REGISTRO)
        # Mismo color en sus matrículas vivas y en las archivadas.
        self.assertEqual(paleta.de_usuario(otra), paleta.de_nombre('Glenda'))
        self.assertEqual(paleta.de_nombre(''), COLOR_SIN_REGISTRO)

    def test_el_color_elegido_por_el_admin_manda(self):
        glenda = User.objects.create_user('glenda', first_name='Glenda', last_name='Ruiz')
        shirley = User.objects.create_user('shirley', first_name='Shirley')
        PerfilUsuario.objects.create(user=glenda, color_registro='morado')
        PerfilUsuario.objects.create(user=shirley, color_registro='blanco')
        paleta = Paleta()
        self.assertEqual(paleta.de_usuario(glenda), HEX['morado'])
        self.assertEqual(paleta.de_usuario(shirley), HEX['blanco'])
        # Las archivadas guardan el nombre: «Glenda Ruiz».
        self.assertEqual(paleta.de_nombre('Glenda Ruiz'), HEX['morado'])
        self.assertEqual(
            color_de_usuario(glenda),
            {'codigo': 'morado', 'nombre': 'Morado', 'hex': HEX['morado'], 'automatico': False},
        )
        self.assertTrue(color_de_usuario(User.objects.create_user('nadie'))['automatico'])

    def test_paleta_tiene_los_colores_del_excel_y_los_nuevos(self):
        self.assertEqual(
            list(HEX),
            ['azul', 'naranja', 'rojo', 'verde', 'rosado', 'morado', 'lila', 'blanco'],
        )


class ColorEnAdminYPerfilTests(TestCase):
    def setUp(self):
        Group.objects.create(name='Asesores').user_set.add(
            User.objects.create_user('melanie', first_name='Melanie', password='x'),
        )
        self.melanie = User.objects.get(username='melanie')
        self.admin = User.objects.create_superuser('jefe', password='x')
        self.url_admin = reverse('admin:auth_user_change', args=[self.melanie.pk])

    def datos_admin(self, color):
        self.client.force_login(self.admin)
        form = self.client.get(self.url_admin).context['adminform'].form
        datos = {}
        for nombre, campo in form.fields.items():
            valor = form.initial.get(nombre, campo.initial)
            if nombre in ('groups', 'user_permissions'):
                datos[nombre] = [g.pk for g in valor] if valor else []
            elif isinstance(valor, bool):
                if valor:
                    datos[nombre] = 'on'
            elif valor is not None and nombre != 'password':
                datos[nombre] = valor.isoformat() if hasattr(valor, 'isoformat') else valor
        datos['color_registro'] = color
        for clave in ('last_login', 'date_joined'):
            fecha = getattr(self.melanie, clave)
            if fecha:
                local = timezone.localtime(fecha)
                datos[f'{clave}_0'] = local.date().isoformat()
                datos[f'{clave}_1'] = local.time().strftime('%H:%M:%S')
                datos.pop(clave, None)
        return datos

    def test_admin_elige_el_color_en_modificar_usuario(self):
        self.client.force_login(self.admin)
        response = self.client.get(self.url_admin)
        self.assertContains(response, 'Color en el Registro Estudiantil')
        self.assertContains(response, 'Automático (ahora: Verde)')
        for nombre in ('Rosado', 'Morado', 'Lila', 'Blanco'):
            self.assertContains(response, nombre)
        response = self.client.post(self.url_admin, self.datos_admin('lila'))
        self.assertEqual(response.status_code, 302, getattr(response, 'context', None) and response.context['adminform'].form.errors)
        self.assertEqual(PerfilUsuario.objects.get(user=self.melanie).color_registro, 'lila')
        self.assertContains(self.client.get(reverse('admin:auth_user_changelist')), 'Lila')

    def test_la_asesora_ve_su_color_en_el_perfil_y_no_puede_cambiarlo(self):
        PerfilUsuario.objects.create(user=self.melanie, color_registro='rosado')
        self.client.force_login(self.melanie)
        url = reverse('academia:comprobante_asesor_detalle', args=[self.melanie.pk])
        response = self.client.get(url)
        self.assertContains(response, 'Tu color en el Registro Estudiantil: Rosado')
        self.assertContains(response, HEX['rosado'])
        self.client.post(url, {'avatar': 'x', 'color_registro': 'azul'})
        self.assertEqual(PerfilUsuario.objects.get(user=self.melanie).color_registro, 'rosado')

    def test_perfil_sin_color_elegido_muestra_el_automatico(self):
        self.client.force_login(self.melanie)
        response = self.client.get(
            reverse('academia:comprobante_asesor_detalle', args=[self.melanie.pk]),
        )
        self.assertContains(response, 'Tu color en el Registro Estudiantil: Verde')
        self.assertContains(response, '(automático)')


class HojaPorDiaTests(RegistroBase):
    def test_sin_parametros_abre_la_hoja_de_hoy_desde_cero(self):
        hoy = timezone.localdate()
        de_hoy = self.matricula('Luis Hoy', hoy, self.kim)
        self.matricula('Eva Ayer', hoy - timedelta(days=1), self.kim)
        self.client.force_login(self.kim)
        response = self.client.get(self.url_hoja)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['es_hoy'])
        self.assertEqual([f['pk'] for f in response.context['filas']], [de_hoy.pk])
        self.assertContains(response, 'se cierra a las 12:00 de la noche')
        hojas = response.context['hojas']
        # Pestañas de la más reciente a la más antigua; la de hoy, activa.
        self.assertEqual([h['fecha'] for h in hojas], sorted((h['fecha'] for h in hojas), reverse=True))
        self.assertTrue(next(h for h in hojas if h['fecha'] == hoy)['activa'])

    def test_hoja_nueva_de_hoy_aparece_aunque_este_vacia(self):
        self.client.force_login(self.kim)
        response = self.client.get(self.url_hoja)
        hoy = timezone.localdate()
        self.assertEqual(response.context['filas'], [])
        self.assertIn(hoy, [h['fecha'] for h in response.context['hojas']])
        self.assertContains(response, 'Hoja nueva: todavía no hay matrículas registradas hoy.')

    def test_hoja_de_un_dia_en_orden_de_registro_con_columnas_del_excel(self):
        segunda = self.matricula('Ana Ruiz', date(2026, 10, 1), self.kim)
        self.matricula('Otro Día', date(2026, 10, 2), self.kim)
        self.client.force_login(self.admin)
        response = self.client.get(self.url_hoja, {'hoja': '2026-10-01'})
        filas = response.context['filas']
        self.assertEqual([f['pk'] for f in filas], [self.m.pk, segunda.pk])
        self.assertEqual([f['numero'] for f in filas], [1, 2])
        for columna in (
            'No.', 'Fecha', 'Registra', 'Nombres del Estudiante', 'Vendedora',
            'Central', 'Programa / Curso', 'Modalidad', 'Jornada',
            'Fecha de inicio', 'Matrícula', 'Valor', 'Forma de Pago', 'Banco',
            'No. de Factura',
        ):
            self.assertContains(response, f'<th>{columna}</th>')
        etiquetas = filas[0]['etiquetas']
        self.assertEqual(etiquetas['registra'], 'Shirley')
        self.assertEqual(etiquetas['vendedora'], 'Glenda')
        self.assertEqual(etiquetas['modalidad'], 'Presencial Guayaquil')
        self.assertEqual(etiquetas['dias'], 'Domingos Intensivos')
        self.assertEqual(etiquetas['inicio'], '11 de octubre')
        # Solo el pago del día de la matrícula; el posterior no se suma.
        self.assertEqual(etiquetas['valor'], '10,00')
        self.assertEqual(etiquetas['forma'], 'Transferencia bancaria')
        self.assertEqual(etiquetas['banco'], 'Pichincha')
        self.assertEqual(filas[0]['color'], HEX['azul'])
        self.assertEqual(filas[1]['color'], HEX['naranja'])
        self.assertContains(response, 'Cerrada')

    def test_asesora_solo_consulta(self):
        self.client.force_login(self.shirley)
        response = self.client.get(self.url_hoja, {'hoja': '2026-10-01'})
        self.assertFalse(response.context['puede_editar'])
        self.assertNotIn('datos_edicion', response.context)
        self.assertNotContains(response, 'registro-datos')
        self.assertContains(response, 'Vista de consulta')

    def test_admin_recibe_los_datos_para_editar_en_la_hoja(self):
        self.client.force_login(self.admin)
        response = self.client.get(self.url_hoja, {'hoja': '2026-10-01'})
        self.assertContains(response, 'id="registro-datos"')
        fila = response.context['filas_edicion'][f'm{self.m.pk}']
        self.assertTrue(fila['pago_editable'])
        self.assertEqual(fila['valores']['valor_pago'], '10.00')
        self.assertEqual(fila['valores']['jornada'], str(self.j_dom.pk))
        self.assertEqual(fila['valores']['sede'], 'Guayaquil')
        # Para mostrar en la confirmación si el valor del curso cambia.
        self.assertEqual((fila['valor_curso'], fila['descuento']), ('90.00', '0.00'))
        datos = response.context['datos_edicion']
        self.assertIn(['venta_presencial', 'Venta Presencial'], [list(c) for c in datos['centrales']])
        self.assertIn(str(self.j_sab.pk), [j['id'] for j in datos['jornadas']])
        curso = next(c for c in datos['cursos'] if c['id'] == str(self.curso.pk))
        self.assertEqual((curso['valor_presencial'], curso['valor_online']), ('90.00', '60.00'))

    def test_las_matriculas_archivadas_siguen_en_su_hoja(self):
        self.archivada(date(2026, 10, 1), original_id=987654)
        # Una copia de una matrícula que sigue viva no se repite.
        self.archivada(date(2026, 10, 1), original_id=self.m.pk, nombre='Duplicada')
        self.client.force_login(self.admin)
        response = self.client.get(self.url_hoja, {'hoja': '2026-10-01'})
        filas = response.context['filas']
        self.assertEqual([f['etiquetas']['nombres'] for f in filas], ['Ana Archivada', 'Chango Zambrano Daniel'])
        archivada = filas[0]
        self.assertTrue(archivada['archivada'])
        self.assertEqual(archivada['etiquetas']['valor'], '10,00')
        self.assertEqual(archivada['etiquetas']['forma'], 'Depósito')
        self.assertEqual(archivada['etiquetas']['central'], 'Central 1')
        self.assertEqual(archivada['color'], HEX['azul'])
        self.assertNotIn(archivada['clave'], response.context['filas_edicion'])
        total = next(h for h in response.context['hojas'] if h['iso'] == '2026-10-01')['total']
        self.assertEqual(total, 2)

    def test_tarjeta_en_el_menu_antes_de_ayuda(self):
        for usuario in (self.admin, self.kim):
            self.client.force_login(usuario)
            contenido = self.client.get(reverse('academia:bienvenida')).content.decode()
            self.assertIn(self.url_hoja, contenido)
            self.assertLess(
                contenido.index('card-title">Registro Estudiantil'),
                contenido.index('card-title">Ayuda'),
            )


class RangoDeFechasTests(RegistroBase):
    def test_sin_rango_solo_aparece_el_selector(self):
        self.client.force_login(self.kim)
        response = self.client.get(self.url_rango)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['filas'], [])
        self.assertContains(response, 'data-date-range-picker')
        self.assertContains(response, 'Ningún rango elegido')
        self.assertNotContains(response, 'aria-label="Hojas por día"')
        self.assertNotContains(response, '<table')

    def test_rango_junta_las_hojas_de_lo_mas_reciente_a_lo_mas_antiguo(self):
        anterior = self.matricula('Ana Ruiz', date(2026, 9, 30), self.kim)
        fuera = self.matricula('Fuera Rango', date(2026, 9, 1), self.kim)
        archivada = self.archivada(date(2026, 9, 30), original_id=555555)
        self.client.force_login(self.admin)
        response = self.client.get(
            self.url_rango, {'fecha_desde': '2026-09-29', 'fecha_hasta': '2026-10-02'},
        )
        claves = [f['clave'] for f in response.context['filas']]
        self.assertEqual(claves[0], f'm{self.m.pk}')
        self.assertEqual(set(claves[1:]), {f'm{anterior.pk}', f'a{archivada.pk}'})
        self.assertNotIn(f'm{fuera.pk}', claves)
        self.assertEqual([f['numero'] for f in response.context['filas']], [1, 2, 3])
        self.assertEqual(response.context['total'], 3)
        self.assertNotContains(response, 'aria-label="Hojas por día"')

    def test_rango_de_un_solo_dia(self):
        self.client.force_login(self.kim)
        response = self.client.get(self.url_rango, {'fecha_desde': '2026-10-01'})
        self.assertEqual([f['pk'] for f in response.context['filas']], [self.m.pk])


class GuardarCeldaTests(RegistroBase):
    def test_solo_el_administrador_puede_guardar(self):
        self.client.force_login(self.shirley)
        response = self.client.post(self.url_guardar, self.datos(nombres='Otro Nombre'))
        self.assertEqual(response.status_code, 403)
        self.assertFalse(response.json()['ok'])
        self.estudiante.refresh_from_db()
        self.assertEqual(self.estudiante.nombres, 'Chango Zambrano Daniel')
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(self.url_guardar).status_code, 405)

    def test_sin_cambios_no_guarda_nada(self):
        response, datos = self.guardar()
        self.assertEqual(response.status_code, 200)
        self.assertTrue(datos['ok'])
        self.assertEqual(datos['cambios'], [])
        self.assertFalse(LogEntry.objects.exists())

    def test_corrige_registra_vendedora_central_nombre_y_factura(self):
        response, datos = self.guardar(
            registrado_por=self.kim.pk, vendedora=self.shirley.pk,
            tipo_registro='venta_presencial',
            nombres='  Chango   Zambrano Daniel Alejandro ', numero_factura='0010443',
        )
        self.assertTrue(datos['ok'])
        self.m.refresh_from_db()
        self.estudiante.refresh_from_db()
        self.assertEqual(self.m.registrado_por, self.kim)
        self.assertEqual(self.m.vendedora, self.shirley)
        self.assertEqual(self.m.tipo_registro, 'venta_presencial')
        self.assertEqual(self.estudiante.nombres, 'Chango Zambrano Daniel Alejandro')
        self.assertEqual(self.m.numero_factura, '0010443')
        self.assertEqual(self.m.factura_realizada, 'si')
        self.assertEqual(self.m.fact_cedula, '0912345678')
        comp = Comprobante.objects.get(matricula=self.m)
        self.assertEqual(comp.vendedora, self.shirley)
        self.assertEqual(comp.tipo_registro, 'venta_presencial')
        self.assertEqual(comp.get_tipo_registro_display(), 'Venta Presencial')
        # La hoja recibe la fila nueva con el color de quien registra.
        fila = datos['fila']
        self.assertEqual(fila['etiquetas']['central'], 'Venta Presencial')
        self.assertEqual(fila['etiquetas']['registra'], 'kim')
        self.assertEqual(fila['color'], HEX['naranja'])
        self.assertEqual(fila['version'], self.m.actualizado.isoformat())
        # Los pagos no se tocan.
        self.assertEqual(self.m.valor_pagado, Decimal('30.00'))
        self.assertEqual(self.m.abonos.count(), 2)
        log = LogEntry.objects.get(object_id=str(self.m.pk))
        self.assertIn('Registra: Shirley Mora → kim', log.change_message)
        self.assertIn('Central: Central IA → Venta Presencial', log.change_message)

    def test_no_se_puede_dejar_vacio_un_dato_que_ya_existe(self):
        response, datos = self.guardar(registrado_por='', vendedora='', tipo_registro='')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(set(datos['campos']), {'registrado_por', 'vendedora', 'tipo_registro'})

    def test_factura_solo_numeros(self):
        response, datos = self.guardar(numero_factura='001-ABC')
        self.assertEqual(response.status_code, 400)
        self.assertTrue(any(m.startswith('No. de factura:') for m in datos['mensajes']))

    def test_otra_persona_guardo_mientras_se_editaba(self):
        self.client.force_login(self.admin)
        datos = self.datos(nombres='Nombre Nuevo')
        datos['version'] = '2000-01-01T00:00:00+00:00'
        response = self.client.post(self.url_guardar, datos)
        self.assertEqual(response.status_code, 400)
        self.assertTrue(response.json()['conflicto'])
        self.estudiante.refresh_from_db()
        self.assertEqual(self.estudiante.nombres, 'Chango Zambrano Daniel')

    def test_matricula_que_ya_no_existe(self):
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse('academia:registro_estudiantil_guardar', args=[999999]), self.datos(),
        )
        self.assertEqual(response.status_code, 404)
        self.assertTrue(response.json()['conflicto'])


class GuardarPagoTests(RegistroBase):
    def test_corrige_valor_forma_y_banco_del_pago_de_la_matricula(self):
        response, datos = self.guardar(valor_pago='8.00', metodo_pago='deposito', banco='guayaquil')
        self.assertTrue(datos['ok'])
        self.reserva.refresh_from_db()
        self.posterior.refresh_from_db()
        self.m.refresh_from_db()
        self.assertEqual(self.reserva.monto, Decimal('8.00'))
        self.assertEqual((self.reserva.metodo, self.reserva.banco), ('deposito', 'guayaquil'))
        self.assertEqual(self.posterior.monto, Decimal('20.00'))
        self.assertEqual(self.m.abonos.count(), 2)
        self.assertEqual(self.m.valor_pagado, Decimal('28.00'))
        self.assertEqual(Comprobante.objects.get(matricula=self.m).pago_abono, Decimal('28.00'))
        self.assertEqual(datos['fila']['etiquetas']['valor'], '8,00')
        self.assertEqual(datos['fila']['etiquetas']['forma'], 'Depósito')

    def test_banco_debe_corresponder_al_metodo(self):
        response, datos = self.guardar(metodo_pago='transferencia', banco='payphone')
        self.assertEqual(response.status_code, 400)
        self.assertIn('banco', datos['campos'])
        response, datos = self.guardar(metodo_pago='efectivo', banco='pichincha')
        self.assertTrue(datos['ok'])
        self.reserva.refresh_from_db()
        self.assertEqual((self.reserva.metodo, self.reserva.banco), ('efectivo', ''))

    def test_pago_no_puede_superar_el_valor_del_curso(self):
        response, datos = self.guardar(valor_pago='80.00')
        self.assertEqual(response.status_code, 400)
        self.assertIn('valor_pago', datos['campos'])
        self.reserva.refresh_from_db()
        self.assertEqual(self.reserva.monto, Decimal('10.00'))

    def test_reserva_es_de_maximo_diez_dolares(self):
        response, datos = self.guardar(valor_pago='10.01')
        self.assertEqual(response.status_code, 400)
        self.assertIn('valor_pago', datos['campos'])
        self.assertIn('máximo $10.00', ' '.join(datos['mensajes']))
        self.reserva.refresh_from_db()
        self.assertEqual(self.reserva.monto, Decimal('10.00'))
        response, datos = self.guardar(valor_pago='0.50')
        self.assertTrue(datos['ok'])

    def test_reserva_antigua_mayor_se_respeta_si_no_cambia(self):
        Abono.objects.filter(pk=self.reserva.pk).update(monto=Decimal('35.00'))
        self.m.recalcular_valor_pagado()
        response, datos = self.guardar(valor_pago='35.00', nombres='Daniel Chango')
        self.assertTrue(datos['ok'])
        response, datos = self.guardar(valor_pago='30.00')
        self.assertEqual(response.status_code, 400)
        self.assertIn('valor_pago', datos['campos'])
        response, datos = self.guardar(valor_pago='9.00')
        self.assertTrue(datos['ok'])

    def test_no_pasa_a_reserva_con_un_pago_mayor_a_diez(self):
        Matricula.objects.filter(pk=self.m.pk).update(
            tipo_matricula='programa_completo', forma_pago='pago_completo',
        )
        self.posterior.delete()
        Abono.objects.filter(pk=self.reserva.pk).update(monto=Decimal('90.00'))
        self.m.recalcular_valor_pagado()
        response, datos = self.guardar(tipo_matricula='reserva_abono', valor_pago='90.00')
        self.assertEqual(response.status_code, 400)
        self.assertIn('tipo_matricula', datos['campos'])
        self.m.refresh_from_db()
        self.assertEqual(self.m.tipo_matricula, 'programa_completo')
        # Bajando a la vez el pago a la reserva, sí se acepta.
        response, datos = self.guardar(tipo_matricula='reserva_abono', valor_pago='10.00')
        self.assertTrue(datos['ok'])

    def test_pago_mixto_no_se_edita_desde_la_hoja(self):
        Abono.objects.filter(pk=self.reserva.pk).update(
            monto_2=Decimal('4.00'), metodo_2='efectivo',
        )
        self.client.force_login(self.admin)
        response = self.client.get(self.url_hoja, {'hoja': '2026-10-01'})
        fila = response.context['filas_edicion'][f'm{self.m.pk}']
        self.assertFalse(fila['pago_editable'])
        self.assertIn('editar_pago=1', fila['enlace_pago'])
        self.assertEqual(fila['etiquetas']['forma'], 'Transferencia bancaria + Efectivo')
        response, datos = self.guardar(valor_pago='99.00', nombres='Daniel Chango')
        self.assertTrue(datos['ok'])
        self.reserva.refresh_from_db()
        self.assertEqual(self.reserva.monto, Decimal('10.00'))

    def test_cambiar_fecha_mueve_el_pago_de_la_matricula_a_otra_hoja(self):
        response, datos = self.guardar(fecha_matricula='2026-10-02')
        self.assertTrue(datos['ok'])
        self.assertEqual(datos['fila']['fecha'], '2026-10-02')
        self.reserva.refresh_from_db()
        self.posterior.refresh_from_db()
        self.assertEqual(self.reserva.fecha, date(2026, 10, 2))
        self.assertEqual(self.posterior.fecha, date(2026, 10, 15))
        response = self.client.get(self.url_hoja, {'hoja': '2026-10-02'})
        self.assertEqual(response.context['filas'][0]['etiquetas']['valor'], '10,00')
        response = self.client.get(self.url_hoja, {'hoja': '2026-10-01'})
        self.assertEqual(response.context['filas'], [])

    def test_fecha_no_puede_pasar_un_pago_posterior(self):
        response, datos = self.guardar(fecha_matricula='2026-10-15')
        self.assertEqual(response.status_code, 400)
        self.assertIn('fecha_matricula', datos['campos'])
        self.reserva.refresh_from_db()
        self.assertEqual(self.reserva.fecha, date(2026, 10, 1))


class GuardarCursoTests(RegistroBase):
    def test_cambia_jornada_y_modalidad_del_mismo_curso(self):
        response, datos = self.guardar(modalidad='online', jornada=self.j_online.pk)
        self.assertTrue(datos['ok'])
        self.m.refresh_from_db()
        self.assertEqual(self.m.jornada, self.j_online)
        self.assertEqual(self.m.modalidad, 'online')
        self.assertEqual(self.m.valor_curso, Decimal('60.00'))
        self.assertEqual(datos['fila']['etiquetas']['modalidad'], 'Virtual Ecuador')
        self.assertEqual(datos['fila']['etiquetas']['inicio'], '20 de octubre')

    def test_la_jornada_debe_ser_del_curso_y_modalidad(self):
        response, datos = self.guardar(jornada=self.j_excel.pk)
        self.assertEqual(response.status_code, 400)
        self.assertIn('jornada', datos['campos'])
        response, datos = self.guardar(modalidad='online')
        self.assertIn('jornada', datos['campos'])

    def test_cambio_de_curso_recalcula_el_valor_y_conserva_los_pagos(self):
        response, datos = self.guardar(curso=self.otro_curso.pk, jornada=self.j_excel.pk)
        self.assertTrue(datos['ok'])
        self.m.refresh_from_db()
        self.assertEqual(self.m.curso, self.otro_curso)
        self.assertEqual(self.m.valor_curso, Decimal('120.00'))
        self.assertEqual(self.m.valor_pagado, Decimal('30.00'))
        self.assertEqual(Comprobante.objects.get(matricula=self.m).curso, self.otro_curso)

    def test_no_cambia_de_curso_si_hay_pagos_de_modulos_que_no_existen(self):
        Abono.objects.create(
            matricula=self.m, fecha=date(2026, 10, 20), monto=Decimal('20.00'),
            metodo='efectivo', tipo_pago='por_modulo', numero_modulo=4,
        )
        response, datos = self.guardar(curso=self.otro_curso.pk, jornada=self.j_excel.pk)
        self.assertEqual(response.status_code, 400)
        self.assertIn('curso', datos['campos'])
        self.m.refresh_from_db()
        self.assertEqual(self.m.curso, self.curso)

    def test_no_cambia_de_curso_con_recuperaciones(self):
        RecuperacionPendiente.objects.create(
            matricula=self.m, numero_modulo=1, fecha_marcada=date(2026, 10, 18),
        )
        response, datos = self.guardar(curso=self.otro_curso.pk, jornada=self.j_excel.pk)
        self.assertEqual(response.status_code, 400)
        self.assertIn('curso', datos['campos'])


class GuardarTipoMatriculaTests(RegistroBase):
    def test_programa_completo_no_cambia_el_valor(self):
        response, datos = self.guardar(tipo_matricula='programa_completo')
        self.assertTrue(datos['ok'])
        self.m.refresh_from_db()
        self.assertEqual(self.m.tipo_matricula, 'programa_completo')
        self.assertEqual(self.m.valor_curso, Decimal('90.00'))

    def test_inscripcion_gratis_descuenta_los_diez_dolares(self):
        response, datos = self.guardar(tipo_matricula='inscripcion_gratis')
        self.assertTrue(datos['ok'])
        self.m.refresh_from_db()
        self.assertEqual(self.m.valor_curso, Decimal('80.00'))
        self.assertEqual(self.m.forma_pago, '')
        self.guardar(tipo_matricula='reserva_abono')
        self.m.refresh_from_db()
        self.assertEqual(self.m.valor_curso, Decimal('90.00'))

    def test_otros_no_se_permite_con_pagos(self):
        response, datos = self.guardar(tipo_matricula='otros')
        self.assertEqual(response.status_code, 400)
        self.assertIn('tipo_matricula', datos['campos'])
        self.m.refresh_from_db()
        self.assertEqual(self.m.valor_curso, Decimal('90.00'))


class VentaPresencialTests(RegistroBase):
    def test_registrar_matricula_acepta_venta_presencial(self):
        from .forms import MatriculaForm
        form = MatriculaForm(prefix='mat')
        self.assertIn(
            ('venta_presencial', 'Venta Presencial'),
            list(form.fields['tipo_registro'].choices),
        )


class FiltrosRegistroTests(RegistroBase):
    def setUp(self):
        super().setUp()
        self.quito = Sede.objects.create(nombre='Quito')
        self.j_quito = JornadaCurso.objects.create(
            curso=self.otro_curso, modalidad='presencial',
            descripcion='sabados_intensivos', fecha_inicio=date(2026, 10, 17),
            sede=self.quito,
        )
        # Mismo día que self.m (01/10): otra asesora, otro curso, con descuento.
        self.otra = Matricula.objects.create(
            estudiante=Estudiante.objects.create(cedula='0955555555', nombres='Ana Quiteña'),
            curso=self.otro_curso, jornada=self.j_quito, modalidad='presencial',
            tipo_matricula='programa_completo', fecha_matricula=date(2026, 10, 1),
            valor_curso=Decimal('120.00'), descuento=Decimal('20.00'),
            registrado_por=self.kim, vendedora=self.kim, tipo_registro='central_1',
        )
        Abono.objects.create(
            matricula=self.otra, fecha=date(2026, 10, 1), monto=Decimal('100.00'),
            metodo='efectivo', tipo_pago='pago_completo',
        )
        self.online = Matricula.objects.create(
            estudiante=Estudiante.objects.create(cedula='0966666666', nombres='Luis Virtual'),
            curso=self.curso, jornada=self.j_online, modalidad='online',
            tipo_matricula='reserva_abono', fecha_matricula=date(2026, 10, 1),
            valor_curso=Decimal('60.00'), registrado_por=self.glenda,
            vendedora=self.glenda,
        )

    def claves(self, **filtros):
        self.client.force_login(self.kim)
        response = self.client.get(self.url_hoja, {'hoja': '2026-10-01', **filtros})
        return [f['clave'] for f in response.context['filas']], response

    def test_sin_filtros_estan_todas_las_del_dia(self):
        claves, response = self.claves()
        self.assertEqual(claves, [f'm{self.m.pk}', f'm{self.otra.pk}', f'm{self.online.pk}'])
        self.assertFalse(response.context['hay_filtros'])
        for texto in (
            'Buscar por cédula, nombre, curso o cédula de factura', 'Todos los cursos',
            '— Primero selecciona un curso —', 'Estado pago: Todos', 'Modalidad: Todas',
            'Ciudad: Todas', 'Sede: Todas', 'Registrador: Todos',
        ):
            self.assertContains(response, texto)
        # En la hoja del día no hay filtro de fecha.
        self.assertNotContains(response, 'data-date-range-picker')

    def test_cada_filtro(self):
        casos = [
            ({'curso': self.otro_curso.pk}, [self.otra]),
            ({'curso': self.curso.pk, 'jornada': self.j_online.pk}, [self.online]),
            ({'descuento': 'si'}, [self.otra]),
            ({'descuento': 'no'}, [self.m, self.online]),
            ({'estado_pago': 'pagado'}, [self.otra]),
            ({'estado_pago': 'pendiente'}, [self.m, self.online]),
            ({'modalidad': 'online'}, [self.online]),
            ({'ciudad': 'Quito'}, [self.otra]),
            ({'sede': f'sede:{self.gye.pk}'}, [self.m]),
            ({'registrador': self.glenda.pk}, [self.online]),
            ({'q': '0955555555'}, [self.otra]),
            ({'q': 'virtual luis'}, [self.online]),
            ({'q': 'Excel Gerencial'}, [self.otra]),
        ]
        for filtros, esperadas in casos:
            with self.subTest(filtros=filtros):
                claves, response = self.claves(**filtros)
                self.assertEqual(claves, [f'm{m.pk}' for m in esperadas])
                self.assertTrue(response.context['hay_filtros'])

    def test_filtros_invalidos_se_ignoran_y_la_jornada_pide_curso(self):
        claves, response = self.claves(
            curso='x', descuento='tal vez', estado_pago='otro', modalidad='marte',
            sede='nada', registrador='abc', jornada=self.j_online.pk,
        )
        self.assertEqual(len(claves), 3)
        self.assertFalse(response.context['hay_filtros'])

    def test_filtros_tambien_en_archivadas(self):
        archivada = self.archivada(date(2026, 10, 1), original_id=444444)
        archivada.registrado_por_nombre = 'Glenda'
        archivada.save()
        claves, _ = self.claves(registrador=self.glenda.pk)
        self.assertEqual(claves, [f'a{archivada.pk}', f'm{self.online.pk}'])
        claves, _ = self.claves(estado_pago='pagado')
        self.assertIn(f'a{archivada.pk}', claves)
        claves, _ = self.claves(ciudad='Guayaquil')
        self.assertIn(f'a{archivada.pk}', claves)

    def test_pestanas_y_rango_conservan_los_filtros(self):
        _, response = self.claves(curso=self.otro_curso.pk)
        self.assertContains(response, f'?hoja=2026-10-01&amp;curso={self.otro_curso.pk}')
        self.assertContains(response, f'/registro-estudiantil/rango/?curso={self.otro_curso.pk}')
        self.assertContains(response, '1 de 3 registros (con filtros)')
        self.assertContains(response, 'href="/registro-estudiantil/?hoja=2026-10-01">Limpiar')

    def test_rango_con_filtros_y_selector_de_fecha(self):
        self.client.force_login(self.admin)
        response = self.client.get(self.url_rango, {
            'fecha_desde': '2026-09-01', 'fecha_hasta': '2026-10-31', 'modalidad': 'online',
        })
        self.assertEqual([f['clave'] for f in response.context['filas']], [f'm{self.online.pk}'])
        self.assertContains(response, 'data-date-range-picker')
        self.assertContains(response, 'Modalidad: Todas')

    def test_excel_con_colores_filtros_y_a4_horizontal(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.client.force_login(self.kim)
        response = self.client.get(
            reverse('academia:registro_estudiantil_export_excel'),
            {'hoja': '2026-10-01', 'descuento': 'no'},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])
        self.assertIn('registro_estudiantil_2026-10-01.xlsx', response['Content-Disposition'])
        ws = load_workbook(BytesIO(response.content)).active
        self.assertIn('Registro Estudiantil · Hoja del', ws['A1'].value)
        self.assertIn('Filtros: Sin descuento', ws['A2'].value)
        self.assertEqual(
            [ws.cell(row=3, column=c).value for c in range(1, 16)],
            ['No.', 'Fecha', 'Registra', 'Nombres del Estudiante', 'Vendedora',
             'Central', 'Programa / Curso', 'Modalidad', 'Jornada', 'Fecha de inicio',
             'Matrícula', 'Valor', 'Forma de Pago', 'Banco', 'No. de Factura'],
        )
        self.assertEqual(ws['D4'].value, 'Chango Zambrano Daniel')
        self.assertEqual(ws['L4'].value, 10.0)
        self.assertEqual(ws['A4'].fill.fgColor.rgb[-6:], HEX['azul'].lstrip('#').upper())
        self.assertEqual(ws['D5'].value, 'Luis Virtual')
        self.assertEqual(ws['A6'].value, 'Total')
        self.assertEqual(ws.auto_filter.ref, 'A3:O5')
        self.assertEqual(ws.page_setup.orientation, 'landscape')
        self.assertEqual(str(ws.page_setup.paperSize), str(ws.PAPERSIZE_A4))
        self.assertEqual(ws.page_setup.fitToWidth, 1)
        self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage)

    def test_excel_de_rango_y_sin_rango(self):
        self.client.force_login(self.kim)
        url = reverse('academia:registro_estudiantil_export_excel')
        response = self.client.get(url, {'fecha_desde': '2026-09-01', 'fecha_hasta': '2026-10-31'})
        self.assertIn('2026-09-01_al_2026-10-31', response['Content-Disposition'])
        response = self.client.get(url, {'fecha_desde': ''})
        self.assertRedirects(response, self.url_rango, fetch_redirect_response=False)

    def test_imprimir_en_a4_horizontal(self):
        self.client.force_login(self.kim)
        response = self.client.get(
            reverse('academia:registro_estudiantil_imprimir'),
            {'hoja': '2026-10-01', 'curso': self.otro_curso.pk},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'size: A4 landscape')
        self.assertContains(response, 'Ana Quiteña')
        self.assertNotContains(response, 'Chango Zambrano Daniel')
        self.assertContains(response, 'Filtros: Curso: Excel Gerencial')
        self.assertContains(response, f'/registro-estudiantil/?hoja=2026-10-01&amp;curso={self.otro_curso.pk}')

    def test_botones_de_excel_e_imprimir_llevan_la_consulta(self):
        _, response = self.claves(modalidad='online')
        self.assertContains(response, '/registro-estudiantil/exportar/excel/?hoja=2026-10-01&amp;modalidad=online')
        self.assertContains(response, '/registro-estudiantil/imprimir/?hoja=2026-10-01&amp;modalidad=online')
